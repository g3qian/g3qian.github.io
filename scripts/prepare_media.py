"""Prepare only the approved public figures. Never reads the course folders.

Requires Pillow and python-pptx. Source decks remain outside this site.
Run with the Codex bundled Python, or a Python environment with these packages.
"""
from pathlib import Path
from io import BytesIO
from PIL import Image, ImageOps, ImageDraw, ImageFont
from pptx import Presentation
import json

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "supporting-material"
OLD = ROOT.parent / "personal-website/images/corrosion-digital-twins"
OUT = ROOT / "images"
MANIFEST = []
FONT = '/System/Library/Fonts/Supplemental/Arial.ttf'
SERIF = '/System/Library/Fonts/Supplemental/Times New Roman.ttf'

def white(im):
    im = im.convert('RGBA')
    bg = Image.new('RGB', im.size, 'white')
    bg.paste(im, mask=im.getchannel('A'))
    return bg

def save(im, target, source, width=2000, photo=False):
    dest = OUT / target
    dest.parent.mkdir(parents=True, exist_ok=True)
    im = white(im)
    im.thumbnail((width, 2400), Image.Resampling.LANCZOS)
    quality = 85 if photo else 90
    im.save(dest, 'WEBP', quality=quality, method=6)
    # Keep the source pixels until compression alone cannot meet the budget.
    while dest.stat().st_size > 390_000:
        if quality > 60:
            quality -= 8
        else:
            im.thumbnail((int(im.width * .85), int(im.height * .85)), Image.Resampling.LANCZOS)
        im.save(dest, 'WEBP', quality=quality, method=6)
    MANIFEST.append(dict(file='images/' + target, source=source, width=im.width, height=im.height, bytes=dest.stat().st_size))
    return im

def shapes(items):
    for item in items:
        if item.shape_type == 6:
            yield from shapes(item.shapes)
        else:
            yield item

def extract(deck, slide, shape):
    prs = Presentation(SOURCE / deck)
    sh = list(shapes(prs.slides[slide-1].shapes))[shape]
    return white(Image.open(BytesIO(sh.image.blob)))

def pair(images, labels=None):
    scaled = []
    for im in images:
        im = white(im)
        im.thumbnail((1000, 800), Image.Resampling.LANCZOS)
        scaled.append(im)
    h = max(im.height for im in scaled) + (65 if labels else 0)
    result = Image.new('RGB', (sum(im.width for im in scaled)+40*(len(scaled)-1), h), 'white')
    draw = ImageDraw.Draw(result)
    x = 0
    for i, im in enumerate(scaled):
        if labels:
            draw.text((x+20, 15), labels[i], font=ImageFont.truetype(FONT, 28), fill='#1F5C7A')
        result.paste(im, (x, 65 if labels else 0))
        x += im.width + 40
    return result

def card(im, target, source):
    im = white(im)
    im.thumbnail((1120, 720), Image.Resampling.LANCZOS)
    result = Image.new('RGB', (1200, 800), 'white')
    result.paste(im, ((1200-im.width)//2, (800-im.height)//2))
    save(result, target, source)

def rotate_panels(deck, slide, shape_ids, labels, angle=90, size=360):
    """Rotated simulation frames placed side by side with short labels."""
    prs = Presentation(SOURCE / deck)
    shs = list(shapes(prs.slides[slide-1].shapes))
    panels = []
    for i in shape_ids:
        im = white(Image.open(BytesIO(shs[i].image.blob))).rotate(angle, expand=True, fillcolor='white')
        box = ImageOps.invert(im.convert('L')).point(lambda v: 255 if v > 25 else 0).getbbox()
        im = im.crop(box); im.thumbnail((size, size), Image.Resampling.LANCZOS); panels.append(im)
    gap, top = 36, 70
    out = Image.new('RGB', (sum(p.width for p in panels) + gap*(len(panels)-1) + 40, max(p.height for p in panels) + top + 20), 'white')
    d = ImageDraw.Draw(out); x = 20
    for p, label in zip(panels, labels):
        d.text((x + p.width/2, top/2), label, font=ImageFont.truetype(FONT, 30), fill='#1C2430', anchor='mm')
        out.paste(p, (x, top)); x += p.width + gap
    return out

bridge = 'EWSHM2026_GuofengQian.pptx'
diss = 'Dissertation Final.pptx'

# Operational modal analysis and temperature effects (EWSHM 2026, public slides)
save(extract(bridge,3,2), 'bridge-oma-temperature/bridge-underside.webp', f'{bridge}, slide 3', width=1600, photo=True)
save(pair([extract(bridge,7,1), extract(bridge,7,2)], ['Frequency Domain Decomposition','Stochastic Subspace Identification']), 'bridge-oma-temperature/oma-methods.webp', f'{bridge}, slide 7; two panels combined, labels added')
for slide, shape, name in [(5,2,'sensor-layout'),(9,2,'frequency-history'),(10,2,'frequency-distributions')]:
    save(extract(bridge,slide,shape), f'bridge-oma-temperature/{name}.webp', f'{bridge}, slide {slide}')
fits = pair([extract(bridge,12,2),extract(bridge,12,3)], ['T1 · Intact bridge','T2 · Damaged bridge'])
save(fits, 'bridge-oma-temperature/bilinear-fits.webp', f'{bridge}, slide 12; panels combined without changing data')
# Retain the earlier twin-bridge overview as the project-navigation image.
overview = white(Image.open(ROOT.parent/'personal-website/images/bridge-monitoring/overview.png'))
overview = overview.resize((1120,round(overview.height*1120/overview.width)), Image.Resampling.LANCZOS)
card(overview, 'bridge-oma-temperature/card.webp', 'Earlier website bridge-monitoring/overview.png; twin-bridge aerial view and T1/T2 photographs')
save(pair([extract(bridge,15,2),extract(bridge,15,3)]), 'bridge-oma-temperature/residuals.webp', f'{bridge}, slide 15')

# Multiphysics corrosion simulation and surrogate modeling
save(Image.open(OLD/'Picture1.png'), 'corrosion-simulation/stress-and-pit-shape.webp', 'Existing real research figure: Picture1.png; composited on white')
save(Image.open(OLD/'surrogate_architecture.png'), 'corrosion-simulation/surrogate-architecture.webp', 'Existing real research figure: surrogate_architecture.png')
save(extract('IWSHM_MU.pptx',3,2), 'corrosion-simulation/field-corrosion.webp', 'IWSHM_MU.pptx, slide 3; photo: Brian Eick, USACE ERDC', photo=True)
dyn = rotate_panels(diss, 25, [6,7,8,3,0], ['Static','4 cycles/yr','8 cycles/yr','16 cycles/yr','32 cycles/yr'])
save(dyn, 'corrosion-simulation/cyclic-loading.webp', f'{diss}, slide 25; five frames rotated to match the slide and labeled')
# Match the physical domains rather than the outer image boxes: the stress
# panel includes a color scale and different margins in the original slide.
def matched_pit_stress():
    morphology = extract(diss,57,4).crop((135,35,866,766))
    stress = extract(diss,57,5)
    domain_size = 500
    scale = domain_size / 505
    morphology = morphology.resize((domain_size,domain_size), Image.Resampling.LANCZOS)
    # Retain the complete color scale and its labels beside the stress domain.
    stress = stress.crop((23,50,742,752))
    stress = stress.resize((round(stress.width*scale),round(stress.height*scale)), Image.Resampling.LANCZOS)
    out = Image.new('RGB',(1400,770),'white')
    out.paste(morphology,(40,150))
    out.paste(stress,(650,150-round(98*scale)))
    draw = ImageDraw.Draw(out)
    for x,label in [(40,'Pit Morphology'),(650,'Stress Concentration')]:
        draw.text((x,40),label,font=ImageFont.truetype(FONT,28),fill='#1F5C7A')
    return out

save(matched_pit_stress(), 'corrosion-simulation/pit-tip-stress.webp', f'{diss}, slide 57; original panels cropped and scaled to equal physical-domain sizes; color scale retained; labels added')
save(extract(diss,58,2), 'corrosion-simulation/crack-initiation-criteria.webp', f'{diss}, slide 58; dissertation Fig. 2.5')
cnn = extract(diss,36,0)
# Trim only blank margins around the original architecture, preserving all labels.
box = ImageOps.invert(cnn.convert('L')).point(lambda v: 255 if v > 25 else 0).getbbox()
cnn = ImageOps.expand(cnn.crop(box), border=30, fill='white')
save(cnn, 'corrosion-simulation/cnn-autoencoder.webp', f'{diss}, slide 36; dissertation Fig. 3.3; blank margins trimmed')
save(extract(diss,35,2), 'corrosion-simulation/cnn-gp-narx.webp', f'{diss}, slide 35; dissertation Fig. 3.1; CNN encoding, recursive GP-NARX forecast, and decoding workflow')
save(extract(diss,46,2), 'corrosion-simulation/surrogate-prediction.webp', f'{diss}, slide 46')
card(Image.open(OLD/'Picture1.png'), 'corrosion-simulation/card.webp', 'Existing Picture1.png')

# Uncertainty propagation and reliability analysis
save(Image.open(OLD/'uncertaintypropagation.png'), 'corrosion-reliability/multiscale-link.webp', 'Existing real research figure: uncertaintypropagation.png; composited on white')
save(Image.open(OLD/'reliability2.png'), 'corrosion-reliability/gate-risk-map.webp', 'Existing real research figure: reliability2.png')
card(Image.open(OLD/'reliability2.png'), 'corrosion-reliability/card.webp', 'Existing reliability2.png')
save(pair([extract(diss,13,3), extract(diss,14,11)], ['Measured · Trend and Noise','Generated Water-Level Histories']), 'corrosion-reliability/water-levels.webp', f'{diss}, slides 13–14; USGS gauge data at Greenup Dam')
save(extract(diss,59,2), 'corrosion-reliability/crack-initiation.webp', f'{diss}, slide 59; Scientific Reports 2022')
save(extract(diss,68,1), 'corrosion-reliability/risk-map-and-gate.webp', f'{diss}, slide 68; MSSP 2023, Fig. 16')

# Bayesian model updating with real measurements
for deck,slide,shape,name in [(diss,74,3,'pit-detection'),('IWSHM_MU.pptx',11,2,'inference-workflow'),('IWSHM_MU.pptx',7,2,'experiment-setup'),('IWSHM_MU.pptx',9,2,'pit-depth-maps'),('IWSHM_MU.pptx',13,2,'experimental-posteriors'),('IWSHM_MU.pptx',15,2,'blind-prediction')]:
    save(extract(deck,slide,shape), f'corrosion-model-updating/{name}.webp', f'{deck}, slide {slide}')
card(extract('IWSHM_MU.pptx',9,2), 'corrosion-model-updating/card.webp', 'IWSHM_MU.pptx, slide 9')

# Generalized damage diagnosis across similar structures
def gate_transfer(gates, W, H, gate_h, fs):
    """Two trimmed gate models with names, roles, and a transfer arrow."""
    trimmed = []
    for g in gates:
        g = white(g); box = ImageOps.invert(g.convert('L')).point(lambda v: 255 if v > 25 else 0).getbbox()
        g = g.crop(box); trimmed.append(g.resize((int(g.width*gate_h/g.height), gate_h), Image.Resampling.LANCZOS))
    im = Image.new('RGB', (W, H), 'white'); d = ImageDraw.Draw(im)
    gap = W//5; x0 = (W - trimmed[0].width - gap - trimmed[1].width)//2; y0 = (H-gate_h)//2 + fs
    xs = [x0, x0 + trimmed[0].width + gap]
    for x, g, name, role in zip(xs, trimmed, ['The Dalles','Bonneville'], ['Source · labeled','Target · unlabeled']):
        im.paste(g, (x, y0))
        d.text((x+g.width/2, y0-fs*2.1), name, font=ImageFont.truetype(FONT, fs), fill='#1C2430', anchor='mm')
        d.text((x+g.width/2, y0-fs*1.05), role, font=ImageFont.truetype(FONT, int(fs*.72)), fill='#5B6573', anchor='mm')
    a0 = xs[0]+trimmed[0].width+gap*0.2; a1 = xs[0]+trimmed[0].width+gap*0.8; ay = y0+gate_h//2
    d.line((a0, ay, a1, ay), fill='#1F5C7A', width=max(3, fs//8))
    d.polygon([(a1+fs*0.5, ay), (a1-fs*0.1, ay-fs*0.35), (a1-fs*0.1, ay+fs*0.35)], fill='#1F5C7A')
    d.text(((a0+a1)/2, ay-fs*0.9), 'transfer', font=ImageFont.truetype(FONT, int(fs*.72)), fill='#1F5C7A', anchor='mm')
    return im

gate_models = [extract('IWSHM_DA.pptx',3,3), extract('IWSHM_DA.pptx',3,4)]
save(gate_transfer(gate_models, 2000, 1000, 760, 44), 'generalized-damage-diagnosis/source-target-gates.webp', 'IWSHM_DA.pptx, slide 3; trimmed, with source and target labels added')
save(gate_transfer(gate_models, 1200, 800, 560, 36), 'generalized-damage-diagnosis/card.webp', 'IWSHM_DA.pptx, slide 3', width=1200)
save(extract('IWSHM_DA.pptx',4,0), 'generalized-damage-diagnosis/dann.webp', 'IWSHM_DA.pptx, slide 4')
save(extract('IWSHM_DA.pptx',10,4), 'generalized-damage-diagnosis/feature-alignment.webp', 'IWSHM_DA.pptx, slide 10')
save(pair([extract('IWSHM_DA.pptx',9,4), extract('IWSHM_DA.pptx',9,3)]), 'generalized-damage-diagnosis/target-predictions.webp', 'IWSHM_DA.pptx, slide 9; the two target-gate (Bonneville) panels')

# Share graphic and touch icon use the site's typography and a real figure.
og = Image.new('RGB',(1200,630),'#FBFAF7');d=ImageDraw.Draw(og)
d.text((65,70),'GUOFENG QIAN',font=ImageFont.truetype(FONT,22),fill='#1F5C7A')
for y,text in [(180,'Reading the Signs'),(260,'of Aging'),(340,'Structures.')]:
    d.text((65,y),text,font=ImageFont.truetype(SERIF,64),fill='#1C2430')
d.text((65,540),'Structural Engineer & Researcher · Vancouver, BC',font=ImageFont.truetype(FONT,21),fill='#5B6573')
photo=ImageOps.fit(extract(bridge,3,2),(400,500),method=Image.Resampling.LANCZOS)
og.paste(photo,(735,65));og.quantize(colors=256).save(OUT/'og-image.png',optimize=True)
icon=Image.new('RGB',(180,180),'#1F5C7A');d=ImageDraw.Draw(icon);d.text((90,85),'GQ',anchor='mm',font=ImageFont.truetype(SERIF,90),fill='#FBFAF7');icon.save(OUT/'apple-touch-icon.png')
# This chart is maintained as SVG; include its provenance without redrawing it.
benchmark = OUT / 'corrosion-reliability/model-runs.svg'
MANIFEST.append(dict(file='images/corrosion-reliability/model-runs.svg',
    source='Qian, Hu, and Todd, MSSP 200 (2023), 110619, Table 2; local reliability benchmark function evaluations; manually maintained SVG',
    width=1200, height=680, bytes=benchmark.stat().st_size))
(ROOT/'data/media-sources.json').write_text(json.dumps(MANIFEST,indent=2)+'\n')
print(f'Prepared {len(MANIFEST)} approved figures; largest: {max(m["bytes"] for m in MANIFEST)//1024} KB')
